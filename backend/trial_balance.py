"""Read a client's prior-year trial balance (Excel) to *propose* their account heads.

Layouts differ wildly (Zoho-style indented sections, Tally-style group rows, coded
charts, audit working papers with comparative columns), so nothing is assumed:
find the header band, find the name column, keep every numeric column under its own
header, and report what was found so the user can confirm or correct it.

Output is evidence and suggestions only. It never sets a client's accounting basis,
never classifies a transaction, and never touches needs_review or
statements.confidence. Local, no external AI.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any, Iterable

from openpyxl import load_workbook

_SCAN_ROWS = 25            # header must be near the top
_NAME_HEADER = re.compile(r"^(account(\s*name)?|particulars|heads?|description|ledger)$", re.I)
_CODE_HEADER = re.compile(r"(account\s*(code|no\.?|number)|^code$)", re.I)
_TYPE_HEADER = re.compile(r"(account\s*type|^type$|^nature$|^category$)", re.I)
_AMOUNT_HEADER = re.compile(r"(debit|credit|balance|opening|closing|net|total|year|ytd|\b20\d\d\b)", re.I)
_TOTAL_ROW = re.compile(r"^(grand\s+)?total\b|^net (profit|loss)\b|^difference\b|^diff\b", re.I)
_PERIOD_TEXT = re.compile(r"\d.*\bto\b.*\d", re.I)   # "1-Jan-25 to 31-Dec-25" titles a block, it is not a column label
_BASIS = re.compile(r"basis\s*:\s*(accrual|cash)\b", re.I)

# Account-name markers that show a ledger is kept on the books (evidence for the basis
# conversation, never a decision).
EVIDENCE_MARKERS: dict[str, re.Pattern[str]] = {
    "receivables": re.compile(r"receivable|debtors", re.I),
    "payables": re.compile(r"payable|creditors", re.I),
    "accruals": re.compile(r"accrued|accured|unbilled", re.I),
    "prepayments": re.compile(r"prepa(id|yment)", re.I),
    "depreciation": re.compile(r"depreciation|accum\.|amorti[sz]ation", re.I),
    "provisions": re.compile(r"provision|end of service", re.I),
}


@dataclass
class Account:
    name: str
    row: int                         # 1-based row in the sheet
    section: str | None = None       # nearest header above (Zoho) or the type column (coded)
    code: str | None = None
    indent: int = 0
    values: dict[str, float] = field(default_factory=dict)   # header text -> number, as in the file


@dataclass
class TrialBalance:
    sheet: str
    entity: str | None = None
    period_text: str | None = None
    stated_basis: str | None = None  # "accrual"/"cash" if the file says so; evidence only
    header_row: int | None = None
    name_column: int | None = None
    amount_columns: dict[int, str] = field(default_factory=dict)
    accounts: list[Account] = field(default_factory=list)
    hierarchy: str = "unknown"       # "sections", "type_column", or "unknown"
    warnings: list[str] = field(default_factory=list)


def _text(cell: Any) -> str:
    return " ".join(str(cell).split()) if isinstance(cell, str) else ""


def _number(cell: Any) -> float | None:
    if isinstance(cell, bool):
        return None
    if isinstance(cell, (int, float)):
        return float(cell)
    if isinstance(cell, str) and cell.strip() == "-":
        return 0.0            # accounting exports print zero as "-"
    return None


def _header_band(rows: list[list[Any]], start: int) -> tuple[int, list[str]]:
    """Combine the header row with up to two following header-only rows (Tally prints
    'Opening / Balance' over two rows); the upper row's text is carried right over blanks."""
    band = [start]
    for nxt in (start + 1, start + 2):
        if nxt < len(rows) and any(_text(c) for c in rows[nxt]) and not any(_number(c) is not None for c in rows[nxt]):
            band.append(nxt)
        else:
            break
    width = max(len(r) for r in rows)
    labels = [""] * width
    for ri in band:
        carried = ""
        for ci in range(width):
            cell = _text(rows[ri][ci]) if ci < len(rows[ri]) else ""
            if _PERIOD_TEXT.search(cell):
                cell = ""
            if cell and ci > 0:          # column 0 is the name header; never carry it over the amounts
                carried = cell
            elif cell:
                carried = ""
            labels[ci] = f"{labels[ci]} {cell if ri == band[-1] or ci == 0 else carried}".strip()
    return band[-1], labels


def _find_header(rows: list[list[Any]]) -> int | None:
    for i, row in enumerate(rows[:_SCAN_ROWS]):
        cells = [_text(c) for c in row]
        has_name = any(_NAME_HEADER.match(c) or c.lower().startswith("account") for c in cells if c)
        amounts = sum(1 for c in cells if c and _AMOUNT_HEADER.search(c))
        if has_name and amounts:
            return i
    # two-row Tally header: amounts appear on the row after the name header
    for i, row in enumerate(rows[:_SCAN_ROWS]):
        if any(_NAME_HEADER.match(_text(c)) for c in row) and i + 1 < len(rows):
            if any(_AMOUNT_HEADER.search(_text(c)) for c in rows[i + 1]):
                return i
    return None


def _title_info(rows: list[list[Any]], upto: int) -> tuple[str | None, str | None, str | None]:
    lines = [ln.strip() for row in rows[:upto] for c in row[:1] if isinstance(c, str) for ln in c.splitlines() if ln.strip()]
    text = " ".join(lines)
    basis = _BASIS.search(text)
    entity = next((ln for ln in lines if not re.match(r"(trial balance|basis|from |as (of|at|on)|\d)", ln, re.I)), None)
    period = next((ln for ln in lines if re.search(r"\b(19|20)\d\d\b", ln)), None)
    return entity, period, basis.group(1).lower() if basis else None


def parse_rows(sheet: str, rows: list[list[Any]]) -> TrialBalance:
    tb = TrialBalance(sheet=sheet)
    rows = [list(r) for r in rows]
    if not rows:
        tb.warnings.append("Sheet is empty.")
        return tb
    head = _find_header(rows)
    if head is None:
        tb.warnings.append("Could not find a header row with an account-name column and amount columns; ask the user to point to them.")
        return tb
    last_header, labels = _header_band(rows, head)
    tb.header_row = head + 1
    tb.entity, tb.period_text, tb.stated_basis = _title_info(rows, head)

    name_col = next((i for i, l in enumerate(labels) if _NAME_HEADER.match(l.split(" ")[0]) and not _CODE_HEADER.search(l) and not _TYPE_HEADER.search(l)), None)
    if name_col is None:
        name_col = next((i for i, l in enumerate(labels) if l.lower().startswith("account") and not _CODE_HEADER.search(l) and not _TYPE_HEADER.search(l)), None)
    if name_col is None:
        tb.warnings.append("No account-name column found.")
        return tb
    tb.name_column = name_col
    code_col = next((i for i, l in enumerate(labels) if _CODE_HEADER.search(l)), None)
    type_col = next((i for i, l in enumerate(labels) if _TYPE_HEADER.search(l)), None)
    tb.amount_columns = {i: l for i, l in enumerate(labels)
                         if i not in (name_col, code_col, type_col) and l and _AMOUNT_HEADER.search(l)}
    seen = [l for l in tb.amount_columns.values()]
    for i, l in list(tb.amount_columns.items()):
        if seen.count(l) > 1:                # same header printed twice (e.g. two years): keep both
            tb.amount_columns[i] = f"{l} [col {i + 1}]"
    if not tb.amount_columns:
        tb.warnings.append("No amount columns recognised.")

    sections: list[tuple[int, str]] = []      # (indent, header text) stack
    saw_sections = False
    for ri in range(last_header + 1, len(rows)):
        row = rows[ri]
        raw = row[name_col] if name_col < len(row) else None
        name = _text(raw)
        if not name or _TOTAL_ROW.match(name):
            continue
        indent = len(raw) - len(raw.lstrip()) if isinstance(raw, str) else 0
        values = {tb.amount_columns[c]: v for c in tb.amount_columns
                  if c < len(row) and (v := _number(row[c])) is not None}
        if not values:                       # name with no numbers: a section/group header
            while sections and sections[-1][0] >= indent:
                sections.pop()
            sections.append((indent, name))
            saw_sections = True
            continue
        while sections and sections[-1][0] >= indent and indent > 0:
            sections.pop()
        section = _text(row[type_col]) if type_col is not None and type_col < len(row) and _text(row[type_col]) else None
        if section is None and sections and indent > sections[-1][0]:
            section = sections[-1][1]
        code = row[code_col] if code_col is not None and code_col < len(row) else None
        tb.accounts.append(Account(name=name, row=ri + 1, section=section,
                                   code=str(code).strip() if code not in (None, "") else None,
                                   indent=indent, values=values))
    if type_col is not None:
        tb.hierarchy = "type_column"
    elif saw_sections and any(a.section for a in tb.accounts):
        tb.hierarchy = "sections"
    elif tb.accounts:
        tb.warnings.append("No indentation, sections or type column: group structure is unknown; ask the user to confirm groups.")
    if not tb.accounts:
        tb.warnings.append("Header found but no account rows read.")
    return tb


def read_workbook(data: bytes) -> list[TrialBalance]:
    """One TrialBalance per visible sheet; a sheet that cannot be read yields a warning, not an error."""
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out: list[TrialBalance] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        out.append(parse_rows(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]))
    return out


def basis_evidence(accounts: Iterable[Account]) -> dict[str, list[str]]:
    """Which kinds of accrual-style ledgers exist (names only). Evidence for the basis
    conversation; the basis itself is only ever what the user confirms."""
    found: dict[str, list[str]] = {}
    for acc in accounts:
        for kind, pat in EVIDENCE_MARKERS.items():
            if pat.search(acc.name):
                found.setdefault(kind, []).append(acc.name)
    return found
